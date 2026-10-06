"""
Scraper Vagas Freelance 99freelas -> Planilha/Notion
Busca vagas de Python/automacao/bot/scraper/dashboard no 99freelas,
extrai dados estruturados e salva em CSV/Excel + opcional Notion.

Nota: Workana removido — bloqueia datacenter IPs via Cloudflare Challenge.
Scraper foca no 99freelas que permite acesso público às vagas.

Stack: Playwright, Pandas, Notion API (opcional)
"""
import os
import asyncio
import logging
from datetime import datetime
from typing import List
from dataclasses import dataclass, asdict

import pandas as pd
from playwright.async_api import async_playwright
from dotenv import load_dotenv

load_dotenv()

logging.basicConfig(format="%(asctime)s - %(levelname)s - %(message)s", level=logging.INFO)
logger = logging.getLogger(__name__)

FREELAS_EMAIL = os.getenv("FREELAS_EMAIL")
FREELAS_PASSWORD = os.getenv("FREELAS_PASSWORD")
NOTION_TOKEN = os.getenv("NOTION_TOKEN")
NOTION_DATABASE_ID = os.getenv("NOTION_DATABASE_ID")

SEARCH_TERMS = [
    "Python automacao", "bot Telegram", "web scraping", "scraper Python",
    "Streamlit dashboard", "FastAPI", "Playwright", "BeautifulSoup",
    "automacao Python", "script Python", "bot Discord", "API integracao",
    "planilha automacao", "Google Sheets API"
]


@dataclass
class Vaga:
    plataforma: str
    titulo: str
    descricao: str
    orcamento: str
    prazo: str
    habilidades: List[str]
    propostas: int
    cliente_avaliacao: float
    cliente_projetos: int
    cliente_pagamentos: str
    url: str
    data_coleta: str
    score: float = 0.0


def calculate_score(vaga: Vaga) -> float:
    score = 50.0
    skills_text = " ".join(vaga.habilidades).lower()
    desc_text = vaga.descricao.lower()
    full_text = f"{skills_text} {desc_text}"

    high_value = ["python", "automacao", "bot", "telegram", "scraping", "playwright",
                  "streamlit", "fastapi", "pandas", "beautifulsoup", "selenium"]
    medium_value = ["api", "integracao", "script", "dashboard", "planilha", "excel",
                    "google sheets", "csv", "json", "sqlite", "postgresql"]

    for kw in high_value:
        if kw in full_text:
            score += 8
    for kw in medium_value:
        if kw in full_text:
            score += 4

    if vaga.propostas < 10:
        score += 10
    elif vaga.propostas < 20:
        score += 5

    if vaga.cliente_avaliacao >= 4.8:
        score += 10
    elif vaga.cliente_avaliacao >= 4.5:
        score += 5

    if vaga.cliente_projetos >= 10:
        score += 5
    elif vaga.cliente_projetos >= 3:
        score += 3

    try:
        orc_num = float(vaga.orcamento.replace("R$", "").replace(".", "").replace(",", ".").strip())
        if orc_num < 300:
            score -= 20
        elif orc_num < 500:
            score -= 10
    except:
        pass

    return max(0, min(100, score))


async def search_99freelas(page, term: str) -> List[Vaga]:
    """Busca vagas no 99freelas (público, sem login necessário)."""
    vagas = []
    try:
        search_url = f"https://www.99freelas.com.br/projects?search={term.replace(' ', '%20')}"
        await page.goto(search_url, wait_until="networkidle")

        # Seletores para cards de projeto
        card_selectors = [
            "[data-testid='project-card']",
            ".project-card",
            ".job-item",
            "article",
            ".project-item",
            "div[class*='card']",
            "div[class*='project']",
            "li",
            ".list-item"
        ]
        cards = []
        for sel in card_selectors:
            try:
                await page.wait_for_selector(sel, timeout=5000)
                cards = await page.query_selector_all(sel)
                if cards:
                    logger.info(f"99freelas: encontrou {len(cards)} cards com seletor: {sel}")
                    break
            except:
                continue

        if not cards:
            logger.warning(f"99freelas: nenhum card encontrado para '{term}'")
            return vagas

        for card in cards[:30]:
            try:
                # Título e URL
                titulo_el = await card.query_selector("a[href*='/project/'], h3 a, h2 a, .title a, [data-testid='title']")
                titulo = await titulo_el.inner_text() if titulo_el else "Sem titulo"
                url = await titulo_el.get_attribute("href") if titulo_el else ""
                if url and not url.startswith("http"):
                    url = f"https://www.99freelas.com.br{url}"

                # Filtra ruído: página de criar projeto, títulos genéricos
                if not url or '/project/' not in url or '/project/new' in url:
                    continue
                if titulo.strip().lower() in ["publique um projeto. é grátis.", "publique um projeto", "novo projeto"]:
                    continue

                # Descrição
                desc_el = await card.query_selector("[data-testid='description'], .description, .project-description, p, .text")
                descricao = await desc_el.inner_text() if desc_el else ""

                # Orçamento
                orc_el = await card.query_selector("[data-testid='budget'], .budget, .value, .price, [class*='budget']")
                orcamento = await orc_el.inner_text() if orc_el else "Nao informado"

                # Prazo
                prazo_el = await card.query_selector("[data-testid='deadline'], .deadline, .time, [class*='deadline']")
                prazo = await prazo_el.inner_text() if prazo_el else "Nao informado"

                # Skills
                skills_els = await card.query_selector_all("[data-testid='skill'], .tag, .skill, [class*='skill'], [class*='tag']")
                habilidades = [await s.inner_text() for s in skills_els]

                # Propostas
                prop_el = await card.query_selector("[data-testid='proposals'], .proposals, .count, [class*='proposal']")
                propostas_text = await prop_el.inner_text() if prop_el else "0"
                propostas = int(''.join(filter(str.isdigit, propostas_text))) if propostas_text else 0

                # Cliente
                cliente_el = await card.query_selector("[data-testid='client'], .client, .user, [class*='client']")
                cliente_text = await cliente_el.inner_text() if cliente_el else ""
                import re
                aval_match = re.search(r"(\d[.,]\d)", cliente_text)
                avaliacao = float(aval_match.group(1).replace(",", ".")) if aval_match else 0
                proj_match = re.search(r"(\d+)\s*projet", cliente_text, re.I)
                projetos = int(proj_match.group(1)) if proj_match else 0

                vaga = Vaga(
                    plataforma="99freelas",
                    titulo=titulo.strip(),
                    descricao=descricao.strip()[:500],
                    orcamento=orcamento.strip(),
                    prazo=prazo.strip(),
                    habilidades=habilidades,
                    propostas=propostas,
                    cliente_avaliacao=avaliacao,
                    cliente_projetos=projetos,
                    cliente_pagamentos="Verificado" if "verificado" in cliente_text.lower() else "Nao verificado",
                    url=url,
                    data_coleta=datetime.now().strftime("%Y-%m-%d %H:%M")
                )
                vaga.score = calculate_score(vaga)
                vagas.append(vaga)

            except Exception as e:
                logger.debug(f"Erro parse card 99freelas: {e}")
                continue

    except Exception as e:
        logger.error(f"Erro busca 99freelas '{term}': {e}")

    return vagas


async def run_scraper() -> List[Vaga]:
    all_vagas = []
    async with async_playwright() as p:
        browser = await p.chromium.launch(headless=True)
        context = await browser.new_context(
            user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
            viewport={"width": 1920, "height": 1080},
            locale="pt-BR",
            timezone_id="America/Sao_Paulo"
        )
        page = await context.new_page()

        logger.info("Buscando 99freelas (publico)...")
        for term in SEARCH_TERMS:
            try:
                logger.info(f"Buscando 99freelas: {term}")
                vagas = await search_99freelas(page, term)
                all_vagas.extend(vagas)
                await asyncio.sleep(3)
            except Exception as e:
                logger.error(f"Erro busca 99freelas '{term}': {e}")

        await browser.close()

    # Remove duplicatas por URL
    seen = set()
    unique_vagas = []
    for v in all_vagas:
        if v.url not in seen:
            seen.add(v.url)
            unique_vagas.append(v)

    # Ordena por score desc
    unique_vagas.sort(key=lambda x: x.score, reverse=True)
    logger.info(f"Total vagas unicas: {len(unique_vagas)}")
    return unique_vagas


def save_to_excel(vagas: List[Vaga], filename: str = "vagas_freelance.xlsx"):
    if not vagas:
        logger.warning("Nenhuma vaga para salvar")
        return
    df = pd.DataFrame([asdict(v) for v in vagas])
    cols_order = ["plataforma", "score", "titulo", "orcamento", "prazo", "propostas",
                  "cliente_avaliacao", "cliente_projetos", "cliente_pagamentos",
                  "habilidades", "descricao", "url", "data_coleta"]
    df = df[cols_order]
    df["habilidades"] = df["habilidades"].apply(lambda x: ", ".join(x) if isinstance(x, list) else x)

    with pd.ExcelWriter(filename, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Vagas")
        ws = writer.sheets["Vagas"]
        for col in ws.columns:
            max_len = max(len(str(cell.value)) for cell in col) if col[0].value else 10
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 50)
        from openpyxl.styles import PatternFill
        green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
        yellow_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=2, max_col=2):
            for cell in row:
                if cell.value and isinstance(cell.value, (int, float)):
                    if cell.value >= 80:
                        for c in ws[cell.row]:
                            c.fill = green_fill
                    elif cell.value >= 60:
                        for c in ws[cell.row]:
                            c.fill = yellow_fill
    logger.info(f"Salvo em {filename}")


def save_to_csv(vagas: List[Vaga], filename: str = "vagas_freelance.csv"):
    if not vagas:
        return
    df = pd.DataFrame([asdict(v) for v in vagas])
    df["habilidades"] = df["habilidades"].apply(lambda x: ", ".join(x) if isinstance(x, list) else x)
    df.to_csv(filename, index=False, encoding="utf-8-sig")
    logger.info(f"Salvo em {filename}")


async def sync_notion(vagas: List[Vaga]):
    if not NOTION_TOKEN or not NOTION_DATABASE_ID:
        logger.info("Notion nao configurado, pulando sync")
        return
    try:
        from notion_client import Client
        notion = Client(auth=NOTION_TOKEN)
        for vaga in vagas[:20]:
            if vaga.score < 60:
                break
            notion.pages.create(
                parent={"database_id": NOTION_DATABASE_ID},
                properties={
                    "Titulo": {"title": [{"text": {"content": vaga.titulo[:100]}}]},
                    "Plataforma": {"select": {"name": vaga.plataforma}},
                    "Score": {"number": round(vaga.score, 1)},
                    "Orcamento": {"rich_text": [{"text": {"content": vaga.orcamento}}]},
                    "Prazo": {"rich_text": [{"text": {"content": vaga.prazo}}]},
                    "Propostas": {"number": vaga.propostas},
                    "Avaliacao Cliente": {"number": vaga.cliente_avaliacao},
                    "Projetos Cliente": {"number": vaga.cliente_projetos},
                    "URL": {"url": vaga.url},
                    "Data Coleta": {"date": {"start": vaga.data_coleta[:10]}},
                }
            )
        logger.info("Sync Notion concluido")
    except Exception as e:
        logger.error(f"Erro sync Notion: {e}")


async def main():
    logger.info("Iniciando scraper de vagas freelance (99freelas)...")
    vagas = await run_scraper()
    if not vagas:
        logger.warning("Nenhuma vaga encontrada")
        return

    timestamp = datetime.now().strftime("%Y%m%d_%H%M")
    save_to_excel(vagas, f"vagas_freelance_{timestamp}.xlsx")
    save_to_csv(vagas, f"vagas_freelance_{timestamp}.csv")
    await sync_notion(vagas)

    print("\n" + "=" * 80)
    print(f"TOP 5 VAGAS (score 0-100) - {datetime.now().strftime('%d/%m/%Y %H:%M')}")
    print("=" * 80)
    for i, v in enumerate(vagas[:5], 1):
        print(f"\n{i}. [{v.plataforma}] Score: {v.score:.1f}/100")
        print(f"   {v.titulo}")
        print(f"   {v.orcamento} | {v.prazo} | {v.propostas} propostas")
        print(f"   Cliente: {v.cliente_avaliacao}/5 ({v.cliente_projetos} projetos) | {v.cliente_pagamentos}")
        print(f"   {v.url}")
        print(f"   Skills: {', '.join(v.habilidades[:8])}")
    print(f"\nTotal: {len(vagas)} vagas unicas | Arquivos salvos com timestamp {timestamp}")


if __name__ == "__main__":
    asyncio.run(main())