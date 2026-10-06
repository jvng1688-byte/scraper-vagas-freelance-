"""
Export utilities for freelance scrapers.
"""
import logging
from typing import List
from datetime import datetime

import pandas as pd

from core.models import Vaga

logger = logging.getLogger(__name__)


def save_to_excel(vagas: List[Vaga], filename: str = "vagas_freelance.xlsx"):
    if not vagas:
        logger.warning("Nenhuma vaga para salvar")
        return
    df = pd.DataFrame([v.to_dict() for v in vagas])
    cols_order = ["plataforma", "score", "titulo", "orcamento", "prazo", "propostas",
                  "cliente_avaliacao", "cliente_projetos", "cliente_pagamentos",
                  "habilidades", "descricao", "url", "data_coleta"]
    df = df[cols_order]
    df["habilidades"] = df["habilidades"].apply(lambda x: ", ".join(x) if isinstance(x, list) else x)

    with pd.ExcelWriter(filename, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Vagas")
        ws = writer.sheets["Vagas"]
        for col in ws.columns:
            max_len = max(len(str(cell.value)) for cell in col) if col[0].value else 10
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 50)
        from openpyxl.styles import PatternFill
        green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
        yellow_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=2, max_col=2):
            for cell in row:
                if cell.value and isinstance(cell.value, (int, float)):
                    if cell.value >= 80:
                        for c in ws[cell.row]:
                            c.fill = green_fill
                    elif cell.value >= 60:
                        for c in ws[cell.row]:
                            c.fill = yellow_fill
    logger.info(f"Salvo em {filename}")


def save_to_csv(vagas: List[Vaga], filename: str = "vagas_freelance.csv"):
    if not vagas:
        return
    df = pd.DataFrame([v.to_dict() for v in vagas])
    df["habilidades"] = df["habilidades"].apply(lambda x: ", ".join(x) if isinstance(x, list) else x)
    df.to_csv(filename, index=False, encoding="utf-8-sig")
    logger.info(f"Salvo em {filename}")


async def sync_notion(vagas: List[Vaga], notion_token: str | None = None, notion_database_id: str | None = None):
    if not notion_token or not notion_database_id:
        logger.info("Notion nao configurado, pulando sync")
        return
    try:
        from notion_client import Client
        notion = Client(auth=notion_token)
        for vaga in vagas[:20]:
            if vaga.score < 60:
                break
            notion.pages.create(
                parent={"database_id": notion_database_id},
                properties={
                    "Titulo": {"title": [{"text": {"content": vaga.titulo[:100]}}]},
                    "Plataforma": {"select": {"name": vaga.plataforma}},
                    "Score": {"number": round(vaga.score, 1)},
                    "Orcamento": {"rich_text": [{"text": {"content": vaga.orcamento}}]},
                    "Prazo": {"rich_text": [{"text": {"content": vaga.prazo}}]},
                    "Propostas": {"number": vaga.propostas},
                    "Avaliacao Cliente": {"number": vaga.cliente_avaliacao},
                    "Projetos Cliente": {"number": vaga.cliente_projetos},
                    "URL": {"url": vaga.url},
                    "Data Coleta": {"date": {"start": vaga.data_coleta[:10]}},
                }
            )
        logger.info("Sync Notion concluido")
    except Exception as e:
        logger.error(f"Erro sync Notion: {e}")